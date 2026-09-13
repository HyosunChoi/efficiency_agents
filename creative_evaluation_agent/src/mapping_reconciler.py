"""Reconciles benchmark/source_assets/ (pre-anonymized 'asset_N' filenames chosen by the
user before this harness ever ran) with benchmark/blind_assets/ (this system's own
sequential Blind_ID, asset_001..asset_0NN, used throughout benchmark/results/*.json).

Built after benchmark/mapping/benchmark_asset_map.xlsx — this system's canonical
source_alias<->Blind_ID record — was accidentally overwritten with a different-schema
file (real creative name mapping) on 2026-09-12. Since the actual files in source_assets/
and blind_assets/ were untouched, the original link is recoverable by exact SHA-256
content match: a source file and its blind copy are byte-identical.

Only unambiguous 1:1 hash matches are auto-confirmed (EXACT_MATCH). Anything else is left
as UNMATCHED/MULTIPLE_MATCH for a human to resolve — this module never guesses which
Blind_ID a source asset became.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.blind_mapper import BLIND_ASSETS_DIR, MAPPING_COLUMNS, MAPPING_PATH, SOURCE_ASSETS_DIR
from src.cache import hash_asset
from src.loader import scan_asset_inventory
from src.schema_guard import SchemaGuardError, assert_safe_to_write

PROJECT_ROOT = Path(__file__).resolve().parent.parent
MAPPING_DIR = PROJECT_ROOT / "benchmark" / "mapping"
REAL_NAME_MAP_PATH = MAPPING_DIR / "real_creative_name_map.xlsx"
CROSSWALK_PATH = MAPPING_DIR / "benchmark_master_crosswalk.xlsx"

EXACT_MATCH = "EXACT_MATCH"
UNMATCHED = "UNMATCHED"
MULTIPLE_MATCH = "MULTIPLE_MATCH"

CROSSWALK_COLUMNS = [
    "Original_Media",
    "Real_Creative_Name",
    "Source_Alias",
    "Blind_ID",
    "Source_File",
    "Blind_File",
    "Source_SHA256",
    "Blind_SHA256",
    "Match_Status",
    "Match_Method",
]

# 이 두 파일은 서로 역할이 다르다 — 한쪽이 실수로 다른 쪽 스키마로 덮어써지는 사고(2026-09-12)가
# 이미 한 번 있었다. 저장 전에 항상 헤더를 검사해서 다른 스키마면 쓰지 않고 에러를 낸다.
CANONICAL_MAP_COLUMNS = tuple(MAPPING_COLUMNS)
REAL_NAME_MAP_EXPECTED_COLUMNS = ("매체", "소재명", "블라인드소재명")


def _read_header(path: Path) -> tuple:
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    header = tuple(c.value for c in next(ws.iter_rows(min_row=1, max_row=1)))
    wb.close()
    return header


@dataclass
class _SourceItem:
    media: str
    alias: str  # asset_N (source_assets에 이미 들어있던 익명 파일명/폴더명)
    asset_type: str
    paths: list[Path]
    sha256: str


@dataclass
class _BlindItem:
    blind_id: str  # asset_00N (이 시스템의 진짜 Blind_ID)
    paths: list[Path]
    sha256: str


def _scan_source_items() -> list[_SourceItem]:
    entries = scan_asset_inventory(SOURCE_ASSETS_DIR)
    return [
        _SourceItem(
            media=e.media_folder_raw,
            alias=e.creative_key_raw,
            asset_type=e.asset_type,
            paths=e.paths,
            sha256=hash_asset(e.paths),
        )
        for e in entries
    ]


def _scan_blind_items() -> list[_BlindItem]:
    items: list[_BlindItem] = []
    if not BLIND_ASSETS_DIR.exists():
        return items
    for entry in sorted(BLIND_ASSETS_DIR.iterdir()):
        if entry.name.startswith("."):
            continue
        if entry.is_dir():
            frame_files = sorted(p for p in entry.iterdir() if p.is_file())
            if not frame_files:
                continue
            items.append(_BlindItem(blind_id=entry.name, paths=frame_files, sha256=hash_asset(frame_files)))
        elif entry.is_file():
            items.append(_BlindItem(blind_id=entry.stem, paths=[entry], sha256=hash_asset([entry])))
    return items


def _file_repr(paths: list[Path], group_name: str) -> str:
    return paths[0].name if len(paths) == 1 else f"[{len(paths)} files: {group_name}/]"


@dataclass
class CrosswalkRow:
    original_media: str
    real_creative_name: str | None
    source_alias: str
    blind_id: str | None
    source_file: str
    blind_file: str | None
    source_sha256: str
    blind_sha256: str | None
    match_status: str
    match_method: str
    source_asset_type: str = ""
    source_paths: tuple = ()
    blind_paths: tuple = ()


def _load_real_name_map() -> dict[tuple[str, str], str]:
    """{(매체, 블라인드소재명): 소재명} — real_creative_name_map.xlsx에서 로드."""
    if not REAL_NAME_MAP_PATH.exists():
        return {}
    header = _read_header(REAL_NAME_MAP_PATH)
    if header != REAL_NAME_MAP_EXPECTED_COLUMNS:
        raise SchemaGuardError(
            f"{REAL_NAME_MAP_PATH.name}의 헤더가 예상과 다릅니다: {header} "
            f"(예상: {REAL_NAME_MAP_EXPECTED_COLUMNS})"
        )
    wb = load_workbook(REAL_NAME_MAP_PATH, read_only=True)
    ws = wb.active
    result = {}
    for media, creative_name, source_alias in ws.iter_rows(min_row=2, values_only=True):
        if source_alias is None:
            continue
        result[(media, source_alias)] = creative_name
    wb.close()
    return result


def reconcile() -> list[CrosswalkRow]:
    source_items = _scan_source_items()
    blind_items = _scan_blind_items()
    real_name_map = _load_real_name_map()

    blind_by_hash: dict[str, list[_BlindItem]] = {}
    for b in blind_items:
        blind_by_hash.setdefault(b.sha256, []).append(b)

    source_by_hash: dict[str, list[_SourceItem]] = {}
    for s in source_items:
        source_by_hash.setdefault(s.sha256, []).append(s)

    rows: list[CrosswalkRow] = []
    for s in source_items:
        candidates = blind_by_hash.get(s.sha256, [])
        same_hash_sources = source_by_hash.get(s.sha256, [])
        real_name = real_name_map.get((s.media, s.alias))
        source_file_repr = _file_repr(s.paths, s.alias)

        if len(candidates) == 1 and len(same_hash_sources) == 1:
            b = candidates[0]
            rows.append(
                CrosswalkRow(
                    original_media=s.media,
                    real_creative_name=real_name,
                    source_alias=s.alias,
                    blind_id=b.blind_id,
                    source_file=source_file_repr,
                    blind_file=_file_repr(b.paths, b.blind_id),
                    source_sha256=s.sha256,
                    blind_sha256=b.sha256,
                    match_status=EXACT_MATCH,
                    match_method="SHA256_EXACT_1TO1",
                    source_asset_type=s.asset_type,
                    source_paths=tuple(s.paths),
                    blind_paths=tuple(b.paths),
                )
            )
        elif len(candidates) == 0:
            rows.append(
                CrosswalkRow(
                    original_media=s.media,
                    real_creative_name=real_name,
                    source_alias=s.alias,
                    blind_id=None,
                    source_file=source_file_repr,
                    blind_file=None,
                    source_sha256=s.sha256,
                    blind_sha256=None,
                    match_status=UNMATCHED,
                    match_method="NONE",
                    source_asset_type=s.asset_type,
                    source_paths=tuple(s.paths),
                )
            )
        else:
            candidate_ids = ",".join(b.blind_id for b in candidates)
            rows.append(
                CrosswalkRow(
                    original_media=s.media,
                    real_creative_name=real_name,
                    source_alias=s.alias,
                    blind_id=candidate_ids,
                    source_file=source_file_repr,
                    blind_file=None,
                    source_sha256=s.sha256,
                    blind_sha256=None,
                    match_status=MULTIPLE_MATCH,
                    match_method=(
                        f"AMBIGUOUS: 이 해시를 공유하는 blind 항목 {len(candidates)}개, "
                        f"source 항목 {len(same_hash_sources)}개"
                    ),
                    source_asset_type=s.asset_type,
                    source_paths=tuple(s.paths),
                )
            )
    return rows


def summarize(rows: list[CrosswalkRow]) -> dict[str, int]:
    counts = {EXACT_MATCH: 0, UNMATCHED: 0, MULTIPLE_MATCH: 0}
    for r in rows:
        counts[r.match_status] = counts.get(r.match_status, 0) + 1
    return counts


def save_crosswalk(rows: list[CrosswalkRow]) -> Path:
    assert_safe_to_write(CROSSWALK_PATH, tuple(CROSSWALK_COLUMNS))
    wb = Workbook()
    ws = wb.active
    ws.title = "Crosswalk"
    ws.append(CROSSWALK_COLUMNS)
    for r in sorted(rows, key=lambda r: (r.original_media, r.source_alias)):
        ws.append(
            [
                r.original_media,
                r.real_creative_name,
                r.source_alias,
                r.blind_id,
                r.source_file,
                r.blind_file,
                r.source_sha256,
                r.blind_sha256,
                r.match_status,
                r.match_method,
            ]
        )
    CROSSWALK_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(CROSSWALK_PATH)
    return CROSSWALK_PATH


def restore_canonical_mapping(rows: list[CrosswalkRow]) -> Path:
    """EXACT_MATCH 항목만 benchmark_asset_map.xlsx(캐노니컬 스키마)로 복원한다.
    UNMATCHED/MULTIPLE_MATCH는 Blind_ID를 확신할 수 없어 포함하지 않는다."""
    assert_safe_to_write(MAPPING_PATH, tuple(CANONICAL_MAP_COLUMNS))

    wb = Workbook()
    ws = wb.active
    ws.title = "Asset_Map"
    ws.append(MAPPING_COLUMNS)

    for r in sorted(rows, key=lambda r: r.blind_id or ""):
        if r.match_status != EXACT_MATCH:
            continue
        original_path = (
            str(SOURCE_ASSETS_DIR / r.original_media / r.source_alias)
            if r.source_asset_type == "carousel"
            else str(r.source_paths[0])
        )
        original_file_name = r.source_alias if r.source_asset_type == "carousel" else r.source_paths[0].name
        blind_path = (
            str(BLIND_ASSETS_DIR / r.blind_id)
            if r.source_asset_type == "carousel"
            else str(r.blind_paths[0])
        )
        ws.append(
            [
                r.blind_id,
                r.original_media,
                r.source_alias,
                original_file_name,
                r.source_asset_type,
                original_path,
                blind_path,
            ]
        )

    MAPPING_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(MAPPING_PATH)
    return MAPPING_PATH
