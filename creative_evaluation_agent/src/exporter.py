"""Excel exporters. 01은 Phase 2 대상으로 완전 구현; 02/03은 Phase 3/6 전 스키마 골격만 (spec §14)."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.evaluator import CreativeEvaluation
from src.winner_interpretation import WinnerInterpretation

HEADER_FILL = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")
HEADER_FONT = Font(bold=True)

ENGINE01_COLUMNS = [
    "Product",
    "Objective",
    "Media",
    "Creative_Name",
    "Creative_ID",
    "Asset_Match_Status",
    "Final_Classification",
    "CTR_Actual",
    "CTR_BM",
    "CPC_Actual",
    "CPC_BM",
    "CPM_Actual",
    "CPM_BM",
    "CVR_Actual",
    "CVR_BM",
    "CPA_Actual",
    "CPA_BM",
    "ROAS_Actual",
    "ROAS_BM",
    "Reliability_Metric",
    "Reliability_Metric_Value",
    "Reliability_Tier",
    "Product_Spend_Share_Pct",
    "Equal_Share_Pct",
    "Scale_Index",
    "Scale_Level",
    "Efficiency_Rank",
    "Delivery_Rank",
    "Allocation_Context",
    "Winner_Type",
    "Primary_Winner",
    "Primary_Winner_Reason",
    "Basis_1",
    "Basis_2",
    "Basis_3",
    "Action",
    "Final_Comment",
    "Status_Flags",
]


def _write_header(ws, columns: list[str]) -> None:
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _autosize(ws, columns: list[str], min_width: int = 12, max_width: int = 45) -> None:
    for idx, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(min_width, min(len(name) + 4, max_width))


def _pct(value: float | None) -> float | None:
    return None if value is None else round(value * 100, 2)


def export_engine01(
    evaluations: list[CreativeEvaluation],
    asset_match_status: dict[str, str],
    winner_interpretations: dict[str, WinnerInterpretation],
    output_path: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Performance_Evaluation"
    _write_header(ws, ENGINE01_COLUMNS)

    for row_idx, ev in enumerate(sorted(evaluations, key=lambda e: (e.product, e.objective, e.creative_id)), start=2):
        bm = ev.bm_kpis
        wi = winner_interpretations.get(ev.creative_id)
        values = [
            ev.product,
            ev.objective,
            ev.media_normalized,
            ev.creative_name_normalized,
            ev.creative_id,
            asset_match_status.get(ev.creative_id, "UNKNOWN"),
            ev.final_classification,
            ev.kpis.ctr,
            bm.ctr if bm else None,
            ev.kpis.cpc,
            bm.cpc if bm else None,
            ev.kpis.cpm,
            bm.cpm if bm else None,
            ev.kpis.cvr,
            bm.cvr if bm else None,
            ev.kpis.cpa,
            bm.cpa if bm else None,
            ev.kpis.roas,
            bm.roas if bm else None,
            ev.reliability_metric_name,
            ev.reliability_metric_value,
            ev.reliability_tier,
            _pct(ev.scale.spend_share) if ev.scale else None,
            _pct(ev.scale.equal_share) if ev.scale else None,
            ev.scale.scale_index if ev.scale else None,
            ev.scale.scale_level if ev.scale else None,
            wi.efficiency_rank if wi else None,
            wi.delivery_rank if wi else None,
            wi.allocation_context if wi else ev.allocation_context,
            wi.winner_type if wi else None,
            wi.primary_winner if wi else None,
            wi.primary_winner_reason if wi else None,
            ev.basis_1,
            ev.basis_2,
            ev.basis_3,
            ev.action,
            ev.final_comment,
            ", ".join(ev.status_flags),
        ]
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    _autosize(ws, ENGINE01_COLUMNS)
    _write_methodology_sheet(wb)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _write_methodology_sheet(wb: Workbook) -> None:
    from src.config_loader import get_evaluation_rules

    rules = get_evaluation_rules()
    ws = wb.create_sheet("Methodology")
    ws.column_dimensions["A"].width = 100

    lines = [
        "Creative Performance Evaluation — Methodology",
        f"생성 시각: {datetime.now().isoformat(timespec='seconds')}",
        "",
        "이 리포트의 모든 최종 판정은 Python deterministic rule engine으로 계산되며,",
        "AI/LLM은 숫자 판정에 관여하지 않는다 (Engine 01 설계 원칙).",
        "",
        "1) Benchmark 그룹: 제품명 x Objective. 반드시 Σ분자/Σ분모로 재계산 (단순 평균 금지).",
        "   CTR BM = ΣClicks/ΣImpressions,  CPC BM = ΣSpend/ΣClicks,  CPM BM = ΣSpend/ΣImpressions*1000",
        "   CVR BM = ΣConversions/ΣClicks,  CPA BM = ΣSpend/ΣConversions,  ROAS BM = ΣRevenue/ΣSpend",
        "",
        "2) Performance 평가 그룹: 제품명 x Objective x 소재명. 위와 동일 공식으로 소재 단위 재계산.",
        "",
        "3) Objective 자동분류 키워드 (config/evaluation_rules.yaml):",
        f"   Conversion: {rules['objective_classification']['conversion_keywords']}",
        f"   Traffic: {rules['objective_classification']['traffic_keywords']}",
        "   위 키워드에 해당하지 않으면 UNKNOWN + CHECK.",
        "",
        "4) Budget Share / Scale (제품명 x 소재명 그룹, Objective 무관):",
        "   Valid Creative = 분석기간 내 Spend > 0",
        "   Equal Share = 100% / 유효 소재 수",
        "   Scale Index = 실제 제품 기준 소재 광고비 비중 / Equal Share",
        "   Scale band: <0.5 Under-delivery, 0.5~<0.8 Low Scale, 0.8~1.2 Normal, >1.2~1.5 High, >1.5 Strong",
        f"   최종 판정 Scale threshold = {rules['scale']['final_classification_threshold']}",
        "",
        "5) Reliability:",
        f"   Conversion: {rules['reliability']['conversion']}",
        f"   Traffic (Clicks 기준): {rules['reliability']['traffic']}",
        "   R0/R1은 확정 Winner/Underperformer 판정을 하지 않는다.",
        "",
        "6) Conversion 최종분류: WINNER / HIDDEN GEM / BUDGET DRAINER / UNDERPERFORMER (R2/R3),",
        "   PROMISING / UNRESOLVED (R0/R1). 회사 Target ROAS/CPA/CPC는 소재평가에 사용하지 않는다.",
        "",
        "7) Traffic 최종분류: TRAFFIC WINNER / TRAFFIC HIDDEN GEM / COST EFFICIENT / HOOK STRONG /",
        "   UNDERPERFORMER (R2/R3), PROMISING / UNRESOLVED (R0/R1). CPC 1차, CTR 2차, CPM은 진단용.",
        "   ※ PROMISING의 'CPC/CTR 유리' 연산자는 명세에 명시되어 있지 않아 config로 분리함: "
        f"현재 값 = {rules.get('traffic_promising_interpretation', {}).get('operator')}"
        " (config/evaluation_rules.yaml > traffic_promising_interpretation).",
        "",
        "8) 분모가 0인 KPI는 계산하지 않고 CHECK:ZERO_DENOMINATOR로 남긴다 (임의 추정 금지).",
        "9) Asset_Match_Status: ASSET_NOT_FOUND / MULTIPLE_ASSET_MATCH / MATCHED / CHECK.",
        "   매체+소재명 완전 일치만 인정하며 의미 추정 매칭은 하지 않는다 (spec §4).",
        "",
        "10) Winner 해석 레이어 (V2, Final_Classification과는 별개 레이어):",
        "   후보 풀 = 이미 Benchmark 이상으로 판정된 소재만 (Conversion: WINNER/HIDDEN GEM,",
        "   Traffic: TRAFFIC WINNER/TRAFFIC HIDDEN GEM). 그 외는 전부 Winner_Type=NONE.",
        "   Efficiency_Rank = 후보 풀 내 순수 효율(ROAS desc, 또는 CPC asc/CTR desc) 순위.",
        "   Delivery_Rank = 후보 풀 내 Spend(예산 소화량) 순위.",
        "   EFFICIENCY_LEADER = 효율 1위, SCALED_WINNER = Scale 조건까지 충족한 소재 중 Delivery 1위,",
        "   BOTH = 두 조건 동일 소재. 새 숫자 threshold 없이 기존 Scale threshold(0.8)만 재사용.",
        "   Primary_Winner: Efficiency Leader==Scaled Winner인 경우만 'Y'로 자동 확정.",
        "   두 후보가 다르면 'CHECK'로 남기고 Primary_Winner_Reason에 근거를 남긴다 — spec V2는",
        "   'ROAS 차이 몇% 이내면 Scale 우선' 같은 고정 threshold를 의도적으로 두지 않았고,",
        "   Human_Override/Human Feedback이 쌓인 뒤 Backtest로 정교화할 예정이다 (config로 분리,",
        "   config/evaluation_rules.yaml > winner_interpretation).",
        "   Allocation_Context(AUTO/MANUAL/UNKNOWN): 광고데이터에 자동배분 여부 컬럼이 없으면",
        "   전부 UNKNOWN — AUTO라고 해서 플랫폼이 '우수소재'로 인증했다는 뜻이 아니며, 단지",
        "   '자동배분 환경에서 상대적으로 높은 Delivery를 확보했다'로만 해석한다 (spec V2 §5).",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line)


def export_engine02_skeleton(output_path: Path) -> None:
    """§6 Engine 02 스키마 골격. Phase 3/5 production 실행 전까지는 컬럼 정의만 제공."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Creative_Content_Analysis"
    columns = [
        "Creative_ID",
        "Media",
        "Creative_Name",
        "Asset_Type",
        "First_3sec_Core_Element",
        "Hook_Type",
        "Product_First_Exposure_Time_Sec",
        "Main_Visual_Focus",
        "Offer_Present",
        "Offer_Type",
        "Offer_First_Exposure_Time_Sec",
        "Demo_Present",
        "Demo_Start_Time_Sec",
        "Actual_Color_Or_Makeup_Result",
        "Before_After_Present",
        "Person_Type",
        "Proof_Type",
        "Subtitle_Prominence",
        "Core_Message",
        "Message_Density",
        "Immediate_Understandability",
        "CTA_Type",
        "CTA_First_Time_Sec",
        "Brand_Product_Name_Exposure",
        "Cut_Pace",
        "Average_Cut_Length_Sec",
        "AI_Predicted_Response",
        "AI_Confidence",
        "Reason_1",
        "Reason_2",
        "Reason_3",
        "Uncertainty_Or_Limitations",
        "Provider",
        "Model",
        "Prompt_Version",
        "Schema_Version",
        "Analysis_Status",
    ]
    _write_header(ws, columns)
    _autosize(ws, columns)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def export_engine03_skeleton(output_path: Path) -> None:
    """§10 Engine 03 스키마 골격. 01/02 결합 로직은 Phase 6 대상이라 여기서는 컬럼만 정의."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Integrated_Creative_Evaluation"
    columns = [
        "Product",
        "Objective",
        "Media",
        "Creative_Name",
        "Creative_ID",
        "01_Final_Classification",
        "01_Winner_Type",
        "01_Primary_Winner",
        "01_Primary_Winner_Reason",
        "Observed_Performance",
        "Observed_Performance_Basis",
        "02_AI_Predicted_Response",
        "02_AI_Confidence",
        "Key_Creative_Observations",
        "AI_Reason_1",
        "AI_Reason_2",
        "AI_Reason_3",
        "Integration_Status",
        "Integrated_Interpretation",
        "Human_Feedback",
        "Human_Override",
        "Final_Comment",
    ]
    _write_header(ws, columns)
    _autosize(ws, columns)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
