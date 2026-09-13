"""Builds benchmark/ground_truth/hidden_ground_truth.xlsx from REAL ad performance data
(spec §9 Observed_Performance), using benchmark_master_crosswalk.xlsx to translate
Blind_ID <-> (Original_Media, Real_Creative_Name).

Never imported by asset_analyzer.py / providers/* — those only ever see CreativeAnalysisInput,
which has no field for any of this. ground_truth.py's read-side lock (both providers' _DONE
markers required) is what actually keeps this file unread until comparison time; this module
only writes it.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.benchmark import compute_benchmarks
from src.config_loader import get_evaluation_rules
from src.evaluator import CreativeEvaluation, evaluate_all
from src.loader import find_input_excel, load_ad_data
from src.mapping_reconciler import CROSSWALK_PATH, EXACT_MATCH
from src.normalizer import build_creative_id, normalize_creative_name, normalize_media
from src.scale import compute_scale

PROJECT_ROOT = Path(__file__).resolve().parent.parent
GROUND_TRUTH_PATH = PROJECT_ROOT / "benchmark" / "ground_truth" / "hidden_ground_truth.xlsx"

GROUND_TRUTH_COLUMNS = [
    "Blind_ID",
    "Observed_Performance",
    "Status",
    "Original_Media",
    "Real_Creative_Name",
    "Product",
    "Objective",
    "KPI_Metric",
    "KPI_Actual",
    "KPI_Benchmark",
    "Gap",
    "Reliability_Tier",
    "Reliability_Metric_Value",
    "01_Final_Classification",
    "Reason",
]


@dataclass
class _CrosswalkEntry:
    original_media: str
    real_creative_name: str | None
    blind_id: str
    match_status: str


def _load_crosswalk() -> list[_CrosswalkEntry]:
    if not CROSSWALK_PATH.exists():
        raise FileNotFoundError(
            f"{CROSSWALK_PATH.name}가 없습니다. 먼저 mapping_reconciler로 crosswalk을 생성하세요."
        )
    wb = load_workbook(CROSSWALK_PATH, read_only=True)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    entries = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        entries.append(
            _CrosswalkEntry(
                original_media=record["Original_Media"],
                real_creative_name=record["Real_Creative_Name"],
                blind_id=record["Blind_ID"],
                match_status=record["Match_Status"],
            )
        )
    wb.close()
    return entries


def _observed_from_evaluation(ev: CreativeEvaluation) -> dict:
    """Conversion은 ROAS, Traffic은 CPC 기준 (config/evaluation_rules.yaml > ground_truth,
    spec §9). AROUND_BENCHMARK는 기본 비활성 — 활성화 시 config의 tolerance_pct 사용."""
    rules = get_evaluation_rules()["ground_truth"]
    around_cfg = rules.get("around_benchmark", {})
    around_enabled = around_cfg.get("enabled", False)
    tolerance = around_cfg.get("tolerance_pct")

    if ev.objective == "Conversion":
        metric_name = "ROAS"
        actual = ev.kpis.roas
        bm = ev.bm_kpis.roas if ev.bm_kpis else None
        higher_is_better = True
    else:
        metric_name = "CPC"
        actual = ev.kpis.cpc
        bm = ev.bm_kpis.cpc if ev.bm_kpis else None
        higher_is_better = False

    if actual is None or bm is None:
        return {
            "kpi_metric": metric_name,
            "kpi_actual": actual,
            "kpi_bm": bm,
            "gap": None,
            "observed_performance": "CHECK",
            "reason": f"{metric_name} actual 또는 Benchmark를 계산할 수 없음 (분모 0 등)",
        }

    gap = actual - bm
    ratio = None
    if bm:
        ratio = (actual - bm) / bm if higher_is_better else (bm - actual) / bm

    if around_enabled and tolerance is not None and ratio is not None and abs(ratio) <= tolerance:
        observed = "AROUND_BENCHMARK"
    elif higher_is_better:
        observed = "ABOVE_BENCHMARK" if actual >= bm else "BELOW_BENCHMARK"
    else:
        observed = "ABOVE_BENCHMARK" if actual <= bm else "BELOW_BENCHMARK"

    return {
        "kpi_metric": metric_name,
        "kpi_actual": actual,
        "kpi_bm": bm,
        "gap": gap,
        "observed_performance": observed,
        "reason": "",
    }


def build_ground_truth(ad_data_path: Path | None = None) -> tuple[Path, dict[str, int]]:
    excel_path = ad_data_path or find_input_excel()
    rows = load_ad_data(excel_path)
    benchmarks = compute_benchmarks(rows)
    scale = compute_scale(rows)
    evaluations = evaluate_all(rows, benchmarks, scale)

    evaluations_by_creative_id: dict[str, list[CreativeEvaluation]] = {}
    for ev in evaluations:
        evaluations_by_creative_id.setdefault(ev.creative_id, []).append(ev)

    crosswalk = _load_crosswalk()

    wb = Workbook()
    ws = wb.active
    ws.title = "Ground_Truth"
    ws.append(GROUND_TRUTH_COLUMNS)

    status_counts: dict[str, int] = {}

    def _emit(row_values: list, status: str) -> None:
        ws.append(row_values)
        status_counts[status] = status_counts.get(status, 0) + 1

    for entry in sorted(crosswalk, key=lambda e: e.blind_id):
        if entry.match_status != EXACT_MATCH:
            _emit(
                [entry.blind_id, "CHECK", "CROSSWALK_NOT_EXACT", entry.original_media,
                 entry.real_creative_name, None, None, None, None, None, None, None, None, None,
                 f"crosswalk Match_Status={entry.match_status} — Blind_ID 매핑 자체가 불확실함"],
                "CROSSWALK_NOT_EXACT",
            )
            continue

        if not entry.real_creative_name:
            _emit(
                [entry.blind_id, "CHECK", "NO_REAL_NAME", entry.original_media,
                 entry.real_creative_name, None, None, None, None, None, None, None, None, None,
                 "crosswalk에 실제 소재명이 없음 (real_creative_name_map.xlsx 확인 필요)"],
                "NO_REAL_NAME",
            )
            continue

        media_norm = normalize_media(entry.original_media)
        creative_norm = normalize_creative_name(entry.real_creative_name)
        creative_id = build_creative_id(media_norm, creative_norm)

        matches = evaluations_by_creative_id.get(creative_id, [])

        if len(matches) == 0:
            _emit(
                [entry.blind_id, "CHECK", "NOT_FOUND_IN_AD_DATA", entry.original_media,
                 entry.real_creative_name, None, None, None, None, None, None, None, None, None,
                 f"실제 광고 데이터에서 creative_id={creative_id}를 찾을 수 없음"],
                "NOT_FOUND_IN_AD_DATA",
            )
            continue

        if len(matches) > 1:
            groups = ", ".join(f"{m.product}/{m.objective}" for m in matches)
            _emit(
                [entry.blind_id, "CHECK", "AMBIGUOUS_MULTIPLE_GROUPS", entry.original_media,
                 entry.real_creative_name, None, None, None, None, None, None, None, None, None,
                 f"creative_id={creative_id}가 여러 Product/Objective에 걸쳐 있어 자동 판정 불가: {groups}"],
                "AMBIGUOUS_MULTIPLE_GROUPS",
            )
            continue

        ev = matches[0]
        result = _observed_from_evaluation(ev)
        status = "OK" if result["observed_performance"] != "CHECK" else "INCOMPLETE_KPI"
        _emit(
            [
                entry.blind_id, result["observed_performance"], status, entry.original_media,
                entry.real_creative_name, ev.product, ev.objective, result["kpi_metric"],
                result["kpi_actual"], result["kpi_bm"], result["gap"], ev.reliability_tier,
                ev.reliability_metric_value, ev.final_classification, result["reason"],
            ],
            status,
        )

    GROUND_TRUTH_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(GROUND_TRUTH_PATH)
    return GROUND_TRUTH_PATH, status_counts
