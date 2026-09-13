"""benchmark/source_assets -> benchmark/blind_assets + benchmark_asset_map.xlsx (spec §9).

Original file/folder names are never changed (spec §15). Re-running this after adding new
source assets keeps existing Blind_IDs stable (never reassigns an ID that's already been
sent to a provider) and only allocates new IDs for newly discovered assets.
"""
from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.loader import AssetEntry, scan_asset_inventory
from src.schema_guard import assert_safe_to_write

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SOURCE_ASSETS_DIR = PROJECT_ROOT / "benchmark" / "source_assets"
BLIND_ASSETS_DIR = PROJECT_ROOT / "benchmark" / "blind_assets"
MAPPING_PATH = PROJECT_ROOT / "benchmark" / "mapping" / "benchmark_asset_map.xlsx"

MAPPING_COLUMNS = [
    "Blind_ID",
    "Original_Media",
    "Original_Creative_Name",
    "Original_File_Name",
    "Asset_Type",
    "Original_Path",
    "Blind_Path",
]


@dataclass
class MappingRow:
    blind_id: str
    original_media: str
    original_creative_name: str
    original_file_name: str
    asset_type: str
    original_path: str
    blind_path: str


def _identity_key(asset: AssetEntry) -> str:
    """Stable identity for an original asset, independent of Blind_ID assignment order."""
    return f"{asset.media_folder_raw}::{asset.creative_key_raw}"


def _load_existing_mapping() -> dict[str, MappingRow]:
    if not MAPPING_PATH.exists():
        return {}
    # 2026-09-12: benchmark_asset_map.xlsx가 실수로 다른 역할의 파일(실제 소재명 매핑표)로
    # 덮어써진 적이 있다. 헤더가 다르면 여기서 명확히 멈춘다 — dict(zip(...))으로 조용히 넘어가면
    # 'Original_Media' 등의 키가 없어 KeyError가 나거나, 최악의 경우 엉뚱한 값을 읽게 된다.
    assert_safe_to_write(MAPPING_PATH, tuple(MAPPING_COLUMNS))

    wb = load_workbook(MAPPING_PATH)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    existing: dict[str, MappingRow] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        key = f"{record['Original_Media']}::{record['Original_Creative_Name']}"
        existing[key] = MappingRow(
            blind_id=record["Blind_ID"],
            original_media=record["Original_Media"],
            original_creative_name=record["Original_Creative_Name"],
            original_file_name=record["Original_File_Name"],
            asset_type=record["Asset_Type"],
            original_path=record["Original_Path"],
            blind_path=record["Blind_Path"],
        )
    return existing


def _next_blind_index(existing: dict[str, MappingRow]) -> int:
    max_idx = 0
    for row in existing.values():
        try:
            idx = int(row.blind_id.replace("asset_", ""))
            max_idx = max(max_idx, idx)
        except ValueError:
            continue
    return max_idx + 1


def _copy_asset_blind(asset: AssetEntry, blind_id: str) -> str:
    BLIND_ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    if asset.asset_type == "carousel":
        dest_dir = BLIND_ASSETS_DIR / blind_id
        dest_dir.mkdir(parents=True, exist_ok=True)
        for i, frame_path in enumerate(sorted(asset.paths, key=lambda p: p.name), start=1):
            shutil.copy2(frame_path, dest_dir / f"{i:02d}{frame_path.suffix.lower()}")
        return str(dest_dir)
    else:
        src = asset.paths[0]
        dest = BLIND_ASSETS_DIR / f"{blind_id}{src.suffix.lower()}"
        shutil.copy2(src, dest)
        return str(dest)


def prepare_blind_assets() -> list[MappingRow]:
    assets = scan_asset_inventory(SOURCE_ASSETS_DIR)
    if not assets:
        logger.warning("benchmark/source_assets에 소재가 없습니다.")
        return []

    existing = _load_existing_mapping()
    next_idx = _next_blind_index(existing)

    all_rows: list[MappingRow] = []
    for asset in sorted(assets, key=lambda a: (a.media_folder_raw, a.creative_key_raw)):
        key = _identity_key(asset)
        if key in existing:
            all_rows.append(existing[key])
            continue

        blind_id = f"asset_{next_idx:03d}"
        next_idx += 1
        blind_path = _copy_asset_blind(asset, blind_id)

        original_path = (
            str((SOURCE_ASSETS_DIR / asset.media_folder_raw / asset.creative_key_raw))
            if asset.asset_type == "carousel"
            else str(asset.paths[0])
        )
        original_file_name = (
            asset.creative_key_raw if asset.asset_type == "carousel" else asset.paths[0].name
        )

        row = MappingRow(
            blind_id=blind_id,
            original_media=asset.media_folder_raw,
            original_creative_name=asset.creative_key_raw,
            original_file_name=original_file_name,
            asset_type=asset.asset_type,
            original_path=original_path,
            blind_path=blind_path,
        )
        all_rows.append(row)
        logger.info("신규 blind asset 등록: %s -> %s", key, blind_id)

    _save_mapping(all_rows)
    return all_rows


def _save_mapping(rows: list[MappingRow]) -> None:
    assert_safe_to_write(MAPPING_PATH, tuple(MAPPING_COLUMNS))
    MAPPING_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Asset_Map"
    ws.append(MAPPING_COLUMNS)
    for row in sorted(rows, key=lambda r: r.blind_id):
        ws.append(
            [
                row.blind_id,
                row.original_media,
                row.original_creative_name,
                row.original_file_name,
                row.asset_type,
                row.original_path,
                row.blind_path,
            ]
        )
    wb.save(MAPPING_PATH)


def load_mapping() -> list[MappingRow]:
    return list(_load_existing_mapping().values())
