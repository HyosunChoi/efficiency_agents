"""Refuses to overwrite an Excel file with a different schema than expected.

Built after benchmark_asset_map.xlsx (this system's canonical source_alias<->Blind_ID
record) was accidentally overwritten with a different-role file (real creative name
mapping) on 2026-09-12 — see mapping_reconciler.py for the recovery. Both blind_mapper.py
and mapping_reconciler.py write to files with fixed, distinct roles; this guard stops one
role's writer from silently clobbering the other's file.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


class SchemaGuardError(Exception):
    """파일을 다른 스키마로 덮어쓰려는 시도를 막을 때 발생시킨다."""


def _read_header(path: Path) -> tuple:
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    header = tuple(c.value for c in next(ws.iter_rows(min_row=1, max_row=1)))
    wb.close()
    return header


def assert_safe_to_write(path: Path, expected_columns: tuple) -> None:
    """파일이 이미 존재하고 헤더가 다르면 쓰기를 거부한다 (역할이 다른 파일을 덮어쓰지 않도록)."""
    if not path.exists():
        return
    existing_header = _read_header(path)
    if existing_header != tuple(expected_columns):
        raise SchemaGuardError(
            f"{path.name}의 기존 헤더가 예상과 다릅니다.\n"
            f"  기존: {existing_header}\n"
            f"  예상: {tuple(expected_columns)}\n"
            "다른 역할의 파일을 실수로 덮어쓸 위험이 있어 저장을 중단합니다. "
            "파일 내용을 확인하고 필요하면 다른 이름으로 옮긴 뒤 다시 실행하세요."
        )
