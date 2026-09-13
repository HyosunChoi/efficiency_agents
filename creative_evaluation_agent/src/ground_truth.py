"""Ground Truth lock (spec §9): comparison 로직은 두 provider 예측이 끝나기 전에는 읽지 못한다."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from src.benchmark_runner import is_provider_done

PROJECT_ROOT = Path(__file__).resolve().parent.parent
GROUND_TRUTH_PATH = PROJECT_ROOT / "benchmark" / "ground_truth" / "hidden_ground_truth.xlsx"

REQUIRED_PROVIDERS = ["openai", "claude"]


class GroundTruthLockedError(Exception):
    pass


@dataclass
class GroundTruthRow:
    blind_id: str
    observed_performance: str
    extra: dict


def is_unlocked() -> bool:
    return all(is_provider_done(p) for p in REQUIRED_PROVIDERS)


def load_ground_truth() -> list[GroundTruthRow]:
    if not is_unlocked():
        missing = [p for p in REQUIRED_PROVIDERS if not is_provider_done(p)]
        raise GroundTruthLockedError(
            "Ground Truth는 아직 unlock되지 않았습니다. 아직 완료되지 않은 provider: "
            + ", ".join(missing)
        )
    if not GROUND_TRUTH_PATH.exists():
        raise FileNotFoundError(f"Ground Truth 파일이 없습니다: {GROUND_TRUTH_PATH}")

    wb = load_workbook(GROUND_TRUTH_PATH)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    rows: list[GroundTruthRow] = []
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, raw_row))
        blind_id = record.pop("Blind_ID", None)
        observed = record.pop("Observed_Performance", None)
        if blind_id is None:
            continue
        rows.append(GroundTruthRow(blind_id=blind_id, observed_performance=observed, extra=record))
    return rows
