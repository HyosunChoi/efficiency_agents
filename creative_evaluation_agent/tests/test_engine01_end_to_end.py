"""Excel -> loader -> matcher -> benchmark -> scale -> evaluator -> exporter, 전체 파이프라인 확인."""
from openpyxl import Workbook, load_workbook

from src.benchmark import compute_benchmarks
from src.evaluator import evaluate_all
from src.exporter import export_engine01
from src.loader import AssetEntry, load_ad_data
from src.matcher import MATCHED, match_all
from src.scale import compute_scale
from src.winner_interpretation import compute_winner_interpretation


def _write_fixture_excel(path):
    wb = Workbook()
    ws = wb.active
    ws.append(["제품명", "목적", "매체", "소재명", "노출수", "클릭수", "광고비", "전환수", "매출"])
    ws.append(["무글틴", "전환", "싱글원_메타", "1+1_0828", 1000, 100, 1000, 6, 1000])
    ws.append(["무글틴", "전환", "싱글원_메타", "발색모음", 1000, 100, 100, 10, 400])
    wb.save(path)


def test_full_engine01_pipeline_produces_excel(tmp_path):
    excel_path = tmp_path / "ad_data.xlsx"
    _write_fixture_excel(excel_path)

    rows = load_ad_data(excel_path)
    assert len(rows) == 2
    assert rows[0].Creative_ID == "싱글원_메타__1+1_0828"

    assets = [
        AssetEntry("싱글원_메타", "1+1_0828", "video", [tmp_path / "1+1_0828.mp4"]),
        AssetEntry("싱글원_메타", "발색모음", "image", [tmp_path / "발색모음.jpg"]),
    ]
    matches = match_all(rows, assets)
    assert all(m.status == MATCHED for m in matches.values())

    benchmarks = compute_benchmarks(rows)
    scale = compute_scale(rows)
    evaluations = evaluate_all(rows, benchmarks, scale)
    assert len(evaluations) == 2
    winner_interpretations = compute_winner_interpretation(evaluations)

    asset_match_status = {row.Creative_ID: matches[row.row_number].status for row in rows}

    out_path = tmp_path / "01_Creative_Performance_Evaluation.xlsx"
    export_engine01(evaluations, asset_match_status, winner_interpretations, out_path)

    assert out_path.exists()
    wb = load_workbook(out_path)
    assert "Performance_Evaluation" in wb.sheetnames
    assert "Methodology" in wb.sheetnames
    ws = wb["Performance_Evaluation"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "Final_Classification" in header
    assert "Winner_Type" in header
    assert "Primary_Winner" in header
    assert "Basis_1" in header
    assert ws.max_row == 3  # header + 2 rows
