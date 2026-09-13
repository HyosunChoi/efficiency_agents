"""Ground truth generation must never guess: unmatched/ambiguous crosswalk entries or
creatives absent from the real ad data must come out as explicit CHECK statuses, and the
Observed_Performance for a resolved entry must be traceable to the exact KPI/BM/Gap used.
"""
from openpyxl import Workbook

from src import ground_truth_builder as gtb
from src import mapping_reconciler
from tests.factories import make_row


def _write_crosswalk(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Original_Media", "Real_Creative_Name", "Source_Alias", "Blind_ID", "Source_File",
            "Blind_File", "Source_SHA256", "Blind_SHA256", "Match_Status", "Match_Method",
        ]
    )
    for r in rows:
        ws.append(r)
    wb.save(path)


def _write_ad_data(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["제품명", "목적", "매체", "소재명", "노출수", "클릭수", "광고비", "전환수", "매출"])
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_exact_match_creative_gets_observed_performance_with_audit_trail(tmp_path, monkeypatch):
    crosswalk_path = tmp_path / "crosswalk.xlsx"
    _write_crosswalk(
        crosswalk_path,
        [["메타", "무글틴_A", "asset_1", "asset_001", "f", "f", "h", "h", "EXACT_MATCH", "SHA256_EXACT_1TO1"]],
    )
    monkeypatch.setattr(mapping_reconciler, "CROSSWALK_PATH", crosswalk_path)
    monkeypatch.setattr(gtb, "CROSSWALK_PATH", crosswalk_path)
    monkeypatch.setattr(gtb, "GROUND_TRUTH_PATH", tmp_path / "hidden_ground_truth.xlsx")

    ad_data_path = tmp_path / "ad.xlsx"
    _write_ad_data(
        ad_data_path,
        [
            ["무글틴", "전환", "메타", "무글틴_A", 1000, 100, 1000, 10, 2000],  # ROAS 2.0
            ["무글틴", "전환", "메타", "무글틴_B", 1000, 100, 1000, 10, 500],   # ROAS 0.5 (BM 대조군)
        ],
    )

    path, counts = gtb.build_ground_truth(ad_data_path)
    assert counts == {"OK": 1}

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    row = dict(zip(header, next(ws.iter_rows(min_row=2, max_row=2, values_only=True))))

    assert row["Blind_ID"] == "asset_001"
    assert row["KPI_Metric"] == "ROAS"
    assert row["KPI_Actual"] == 2.0
    # BM = (2000+500)/(1000+1000) = 1.25
    assert row["KPI_Benchmark"] == 1.25
    assert row["Observed_Performance"] == "ABOVE_BENCHMARK"
    assert row["Status"] == "OK"


def test_non_exact_crosswalk_entry_is_check_not_guessed(tmp_path, monkeypatch):
    crosswalk_path = tmp_path / "crosswalk.xlsx"
    _write_crosswalk(
        crosswalk_path,
        [["메타", "무글틴_A", "asset_1", "asset_042,asset_050", "f", None, "h", None, "MULTIPLE_MATCH", "AMBIGUOUS"]],
    )
    monkeypatch.setattr(gtb, "CROSSWALK_PATH", crosswalk_path)
    monkeypatch.setattr(gtb, "GROUND_TRUTH_PATH", tmp_path / "hidden_ground_truth.xlsx")

    ad_data_path = tmp_path / "ad.xlsx"
    _write_ad_data(ad_data_path, [["무글틴", "전환", "메타", "무글틴_A", 1000, 100, 1000, 10, 2000]])

    path, counts = gtb.build_ground_truth(ad_data_path)
    assert counts == {"CROSSWALK_NOT_EXACT": 1}


def test_creative_not_found_in_real_ad_data_is_check(tmp_path, monkeypatch):
    crosswalk_path = tmp_path / "crosswalk.xlsx"
    _write_crosswalk(
        crosswalk_path,
        [["메타", "존재안함", "asset_1", "asset_001", "f", "f", "h", "h", "EXACT_MATCH", "SHA256_EXACT_1TO1"]],
    )
    monkeypatch.setattr(gtb, "CROSSWALK_PATH", crosswalk_path)
    monkeypatch.setattr(gtb, "GROUND_TRUTH_PATH", tmp_path / "hidden_ground_truth.xlsx")

    ad_data_path = tmp_path / "ad.xlsx"
    _write_ad_data(ad_data_path, [["무글틴", "전환", "메타", "다른소재", 1000, 100, 1000, 10, 2000]])

    path, counts = gtb.build_ground_truth(ad_data_path)
    assert counts == {"NOT_FOUND_IN_AD_DATA": 1}


def test_creative_in_multiple_objective_groups_is_ambiguous_check(tmp_path, monkeypatch):
    crosswalk_path = tmp_path / "crosswalk.xlsx"
    _write_crosswalk(
        crosswalk_path,
        [["메타", "무글틴_A", "asset_1", "asset_001", "f", "f", "h", "h", "EXACT_MATCH", "SHA256_EXACT_1TO1"]],
    )
    monkeypatch.setattr(gtb, "CROSSWALK_PATH", crosswalk_path)
    monkeypatch.setattr(gtb, "GROUND_TRUTH_PATH", tmp_path / "hidden_ground_truth.xlsx")

    ad_data_path = tmp_path / "ad.xlsx"
    # 같은 creative_id가 Conversion과 Traffic 두 그룹에 다 걸림
    _write_ad_data(
        ad_data_path,
        [
            ["무글틴", "전환", "메타", "무글틴_A", 1000, 100, 1000, 10, 2000],
            ["무글틴", "유입", "메타", "무글틴_A", 1000, 100, 1000, 0, 0],
        ],
    )

    path, counts = gtb.build_ground_truth(ad_data_path)
    assert counts == {"AMBIGUOUS_MULTIPLE_GROUPS": 1}
